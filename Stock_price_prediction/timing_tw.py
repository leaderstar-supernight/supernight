"""WHID stage 2 orchestration. No orders, no edits to WHID, no embedded secrets."""
from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd

from timing_data import (VERSION, Provider, candidates, load_config, number, read_holdings)
from timing_rules import signals, latest_assessment, compare_strategies, market_sentiment, news_sentiment
from timing_ai import run_ai, summarize_model_metrics
from timing_timesfm import TimesFMPathRunner, summary_columns

SIMPLE_COLUMNS=['代號','訊號日期','現價','短期趨勢','中期趨勢','長期趨勢','相對大盤動能','籌碼','籌碼K線','市場情緒','個股情緒','融資壓力','新聞情緒','風險','未持有建議','持有情境建議',
    'ATR風險參考價','支撐參考','壓力參考','AI同區間最佳模型','AI上行先觸機率','AI下行先觸機率','AI盤整機率',
    *summary_columns('TW'),'AI狀態','資料狀態']

SUMMARY_COLUMNS=['代號','產業','現價','本益比（TTM）','歷史PE百分位_1Y','歷史PE百分位_3Y','歷史PE百分位_5Y',
    '成長分數','品質分數','巴菲特檢核分數','基本面分數','基本面評價星等','基本面投資評價','價格投資建議','盈餘趨勢評價','現金獲利品質評價','財務安全性評價',
    '分析師共識（強買/買進/持有/賣出/強賣）','分析師目標價最低','分析師目標價中位數','分析師目標價平均','分析師目標價最高','目標價依據',
    '營收逐季趨勢','EPS逐季趨勢','OCF逐季趨勢','短期趨勢','中期趨勢','長期趨勢','相對大盤動能',
    '籌碼','籌碼K線','市場情緒','個股情緒','融資壓力','新聞情緒',
    '未持有建議','持有情境建議','ATR風險參考價','支撐參考','壓力參考','AI同區間最佳模型','AI上行先觸機率','AI下行先觸機率','AI盤整機率',
    'TimesFM10日路徑判讀','TimesFM第10日預測價','TimesFM10日預測漲跌幅','TimesFM預估觸及日']

SUMMARY_GROUPS=[
    ('股票代號＆產業別',1,2,'1F4E78'),('現價位於歷史PE地位',3,7,'5B9BD5'),
    ('基本面分析＆買價建議',8,17,'70AD47'),('分析師＆投信買賣建議＆建議價格',18,23,'C55A11'),
    ('最近3季營收、EPS、可操作現金流趨勢',24,26,'8064A2'),('技術分析',27,30,'2F75B5'),
    ('籌碼面分析＆大盤情緒／個股情緒',31,36,'548235'),
    ('AI模型：未持有／持有情境與10日觸及機率',37,45,'BF9000'),
    ('時間序模型：股價預測路徑',46,49,'7030A0')]


def safe_json(value):
    if isinstance(value,dict): return {str(k):safe_json(v) for k,v in value.items()}
    if isinstance(value,(list,tuple)): return [safe_json(v) for v in value]
    if isinstance(value,(pd.Timestamp,datetime)): return value.isoformat()
    if isinstance(value,np.integer): return int(value)
    if isinstance(value,np.bool_): return bool(value)
    if isinstance(value,(float,np.floating)): return float(value) if np.isfinite(value) else None
    if value is pd.NA or value is pd.NaT: return None
    return value


def _sheet(frame):
    if frame is None or frame.empty: return pd.DataFrame({'狀態':['本次無資料；請查看DataStatus與AIStatus']})
    f=frame.copy()
    for col in f:
        if f[col].dtype==object:
            f[col]=f[col].map(lambda v: json.dumps(safe_json(v),ensure_ascii=False) if isinstance(v,(dict,list)) else v)
            # External strings must never become executable spreadsheet formulas.
            f[col]=f[col].map(lambda v: "'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v)
    return f


def _remove_consumed_step1_simple(provenance, report_root):
    """Remove only the exact Step-1 simple report consumed by this successful run."""
    source = Path(str((provenance or {}).get('來源') or ''))
    if not source.is_absolute():
        return None
    expected_parent = (report_root/'簡化版'/'TW').resolve()
    try:
        source = source.resolve()
    except OSError:
        return None
    if (source.parent == expected_parent
            and re.fullmatch(r'TW_20\d{6}_\d{6}_Step1_Report\.xlsx', source.name)
            and source.is_file()):
        source.unlink()
        return str(source)
    return None


def _combined_simple_report(result):
    """Join Step 1 then Step 2 by ticker, keeping one copy of repeated fields."""
    step1=result['candidates'].copy()
    step2=result['simple'].copy()
    if '代號' not in step1 or '代號' not in step2:
        raise ValueError('簡化版整合需要Step1與Step2都包含代號')
    if step1['代號'].duplicated().any() or step2['代號'].duplicated().any():
        raise ValueError('簡化版整合不能包含重複股票代號')
    # Step1的現價是基本面估值採用價；同一列不再重複Step2現價。
    repeated=set(step1.columns).intersection(step2.columns)-{'代號'}
    if '現價' in step1.columns: repeated.add('現價')
    step2=step2.drop(columns=[column for column in repeated if column in step2],errors='ignore')
    return step1.merge(step2,on='代號',how='outer',sort=False,validate='one_to_one')


def _write_summary_report(path, combined):
    from openpyxl.styles import Alignment,Font,PatternFill
    from openpyxl.utils import get_column_letter
    summary=combined.reindex(columns=SUMMARY_COLUMNS).copy()
    with pd.ExcelWriter(path,engine='openpyxl') as writer:
        out=_sheet(summary)
        out.to_excel(writer,sheet_name='Report',index=False,startrow=1)
        ws=writer.sheets['Report']
        for title,start,end,color in SUMMARY_GROUPS:
            ws.merge_cells(start_row=1,start_column=start,end_row=1,end_column=end)
            cell=ws.cell(1,start,title)
            cell.font=Font(color='FFFFFF',bold=True,size=11)
            cell.fill=PatternFill('solid',fgColor=color)
            cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        for cell in ws[2]:
            cell.font=Font(color='FFFFFF',bold=True)
            cell.fill=PatternFill('solid',fgColor='17365D')
            cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        ws.row_dimensions[1].height=34;ws.row_dimensions[2].height=52
        ws.freeze_panes='C3'
        ws.auto_filter.ref=f'A2:{get_column_letter(ws.max_column)}{ws.max_row}'
        for index,column in enumerate(out.columns,1):
            width=16
            if any(word in str(column) for word in ['建議','評價','趨勢','情緒','判讀','模型','共識']): width=22
            if len(str(column))>=18: width=28
            ws.column_dimensions[get_column_letter(index)].width=width
            for cell in ws.iter_cols(min_col=index,max_col=index,min_row=3):
                for value_cell in cell:
                    if not isinstance(value_cell.value,(int,float)): continue
                    if '歷史PE百分位' in str(column): value_cell.number_format='0.0"%"'
                    elif any(word in str(column) for word in ['機率','漲跌幅']): value_cell.number_format='0.0%'
                    elif '分數' in str(column): value_cell.number_format='0'
                    else: value_cell.number_format='#,##0.000'
    return summary


def export_reports(result, cfg):
    now=datetime.now()
    stamp=now.strftime('%Y%m%d_%H%M%S')
    run_id=f'TW_{stamp}_Step2'
    report_root=Path(cfg['output']['root'])
    system_root=Path(cfg['output']['system_root'])
    paths={'詳細版':report_root/'詳細版'/'TW',
           '簡化版':report_root/'簡化版'/'TW',
           '摘要':report_root/'摘要'/'TW',
           '研究資料':system_root/'research'/'TW'}
    for p in paths.values(): p.mkdir(parents=True,exist_ok=True)
    simple=paths['簡化版']/f'TW_{stamp}_Combined_Report.xlsx'
    summary_path=paths['摘要']/f'TW_{stamp}_Summary_Report.xlsx'
    detailed=paths['詳細版']/f'{run_id}_Report.xlsx'
    rules=[{'分類':g,'設定':k,'值':v} for g in ['rules','sentiment','ai','timesfm','backtest'] for k,v in cfg[g].items()]
    rules.extend([
        {'分類':'說明','設定':'趨勢分層','值':'短期5/20/3、中期20/60/5、長期120/240/20；只有中期趨勢參與現行進出場與回測'},
        {'分類':'說明','設定':'籌碼K線','值':'法人、自營商、融資融券、借券與可用的券商分點形成獨立結論；目前未參與交易建議'},
        {'分類':'說明','設定':'市場／個股／新聞情緒','值':'透明規則的研究參考，保存每日輸入供後續驗證；目前不參與交易建議、回測或AI特徵'},
        {'分類':'限制','設定':'融資維持率','值':'公開資料無法還原投資人帳戶擔保品與負債，因此只顯示個股與大盤融資壓力，不冒充實際維持率'},
        {'分類':'限制','設定':'券商分點','值':'FinMind sponsor限定；無權限時自動降級，主要買超分點均價不是CMoney主力成本'},
        {'分類':'說明','設定':'AI目標','值':'未來10個交易日先觸及上方1.5倍ATR、下方1.0倍ATR，或期間內兩者皆未觸及；同日雙觸採下方'},
        {'分類':'說明','設定':'Momentum Benchmark','值':'TAIEX 20/60/120日報酬、個股相對TAIEX報酬與60日風險調整動能；同時作為所有AI的新特徵及獨立比較模型'},
        {'分類':'說明','設定':'AI模型比較','值':'AIComparison以相同126交易日final_test，綜合低LogLoss、低Brier與高MacroF1排名；AIModelSummary依樣本數彙總全部股票'},
        {'分類':'限制','設定':'AI決策','值':'所有AI機率與Momentum Benchmark只作研究參考，未用於交易建議；LSTM與Momentum Benchmark不加入正式集成'},
        {'分類':'說明','設定':'TimesFM 10日路徑','值':'TimesFM 3.0依還原收盤價預測未來10個交易日價格路徑；點路徑判讀先觸方向，0.1至0.9分位路徑另估三分類情境機率'},
        {'分類':'限制','設定':'TimesFM用途','值':'個人非商業研究；不參與AI集成、分數、趨勢、籌碼、風險或任何交易建議'},
        {'分類':'限制','設定':'TimesFM預估日期','值':'路徑以10個交易日序為準；日期欄只排除週末，未排除台股休市日'},
        {'分類':'限制','設定':'回測','值':'固定候選名單的單檔研究，不是WHID歷史選股或投資組合績效'},
        {'分類':'限制','設定':'成交','值':'前日訊號下一日開盤；同日雙觸價採停損優先；零量/一價日不成交'},
        {'分類':'限制','設定':'價格','值':'還原價、分數股與股息再投資研究假設；未模擬實際張數/最低手續費'},
        {'分類':'限制','設定':'成本','值':'買賣費率、賣出稅率及滑價為設定假設，請依商品及券商確認'},
        {'分類':'限制','設定':'資訊時間','值':'台灣20:00前不採當日日線；籌碼再延遲至少1交易日'},
        {'分類':'限制','設定':'風險參考價','值':'目前收盤價減ATR倍數；不是保證停損成交價，也不是個人化部位建議'},
    ])
    sheets={'Report':result['detail'],'DataStatus':result['data_status'],'Backtest':result['backtest'],
        'Trades':result['trades'],'Equity':result['equity'],'AILatest':result['ai_latest'],
        'AIMetrics':result['ai_metrics'],'AIComparison':result['ai_comparison'],
        'AIModelSummary':result['ai_model_summary'],'AICalibration':result['ai_calibration'],'AIStatus':result['ai_status'],
        'TimesFMPath':result.get('timesfm_path',pd.DataFrame()),
        'News':result['news'],'MarketContext':result['market_context'],
        'Rules':pd.DataFrame(rules),'CandidateSource':pd.DataFrame([result['provenance']]),
        'Candidates':result['candidates']}
    # This is the program's reusable export function, not a change to any existing workbook.
    from openpyxl.styles import Font,PatternFill,Alignment
    combined=_combined_simple_report(result)
    simple_sheets={'Report':combined}
    for path,book in [(simple,simple_sheets),(detailed,sheets)]:
        with pd.ExcelWriter(path,engine='openpyxl') as writer:
            for name,df in book.items():
                out=_sheet(df)
                out.to_excel(writer,sheet_name=name,index=False)
                ws=writer.sheets[name]; ws.freeze_panes='C2'; ws.auto_filter.ref=ws.dimensions
                for cell in ws[1]:
                    cell.font=Font(color='FFFFFF',bold=True); cell.fill=PatternFill('solid',fgColor='17365D')
                    cell.alignment=Alignment(wrap_text=True,vertical='center')
                ws.row_dimensions[1].height=32
                for j,(col,cells) in enumerate(zip(out.columns,ws.iter_cols()),1):
                    from openpyxl.utils import get_column_letter
                    ws.column_dimensions[get_column_letter(j)].width=min(42,max(14,len(str(col))*1.7))
                    for cell in list(cells)[1:]:
                        if isinstance(cell.value,(int,float)):
                            cell.number_format='0.0%' if any(s in str(col).lower() for s in ['probability','return','drawdown','win_rate','比例','機率','報酬率','乖離','漲跌幅']) else '#,##0.000'
    _write_summary_report(summary_path,combined)
    snapshot=paths['研究資料']/(run_id+'_snapshot.json')
    snapshot.write_text(json.dumps(safe_json({'version':VERSION,'created_utc':datetime.now(timezone.utc),
        'config':cfg,'provenance':result['provenance'],'rows':result['detail'].to_dict('records')}),ensure_ascii=False,indent=2),encoding='utf-8')
    # Preserve per-day out-of-sample outputs without bloating the simplified report.
    if not result['ai_predictions'].empty:
        result['ai_predictions'].to_csv(paths['研究資料']/(run_id+'_ai_oos.csv'),index=False,encoding='utf-8-sig')
    result['signals'].to_csv(paths['研究資料']/(run_id+'_signals.csv'),index=False,encoding='utf-8-sig')
    if not result['news'].empty:
        result['news'].to_csv(paths['研究資料']/(run_id+'_news.csv'),index=False,encoding='utf-8-sig')
    removed_step1=_remove_consumed_step1_simple(result.get('provenance'), report_root)
    output={'簡化版':str(simple),'摘要':str(summary_path),'詳細版':str(detailed),'快照':str(snapshot)}
    if removed_step1: output['已整併並移除Step1簡化版']=removed_step1
    return output


def analyze_bundle(ticker, candidate, raw, adjusted, chips, meta, cfg, provenance, asof, holding=None,
                   market_summary=None, news_summary=None, benchmark=None, timesfm_runner=None):
    x=signals(raw,adjusted,chips,cfg)
    row=latest_assessment(ticker,raw,adjusted,x,cfg,provenance,asof,holding)
    for col,value in candidate.items():
        if col!='代號': row['WHID_'+col]=value
    row.update(meta)
    row.update(market_summary or {})
    row.update(news_summary or {})
    stale=(pd.Timestamp(asof).normalize()-raw.index[-1]).days>cfg['data']['max_price_age_days']
    if stale:
        ai={'latest':pd.DataFrame(),'metrics':pd.DataFrame(),'calibration':pd.DataFrame(),
            'comparison':pd.DataFrame(),'momentum_snapshot':{'available':False,'label':'資料不足'},
            'predictions':pd.DataFrame(),'errors':[],'status':'行情過期，AI未執行'}
    else: ai=run_ai(adjusted,x,cfg,benchmark)
    timesfm_runner=timesfm_runner or TimesFMPathRunner(cfg,'TW')
    if stale: timesfm_summary,timesfm_path=timesfm_runner.unavailable('行情過期，TimesFM未執行')
    else: timesfm_summary,timesfm_path=timesfm_runner.forecast(ticker,raw,adjusted)
    row.update(timesfm_summary)
    row['AI狀態']=ai['status']
    momentum=ai.get('momentum_snapshot',{})
    row['動能基準']=momentum.get('benchmark',cfg['data'].get('benchmark','TAIEX'))
    row['相對大盤動能']=momentum.get('label','資料不足')
    row['相對大盤20日報酬差']=momentum.get('relative_return_20',np.nan)
    row['相對大盤60日報酬差']=momentum.get('relative_return_60',np.nan)
    row['相對大盤120日報酬差']=momentum.get('relative_return_120',np.nan)
    row['大盤20日報酬率']=momentum.get('benchmark_return_20',np.nan)
    row['大盤60日報酬率']=momentum.get('benchmark_return_60',np.nan)
    row['大盤120日報酬率']=momentum.get('benchmark_return_120',np.nan)
    row['60日風險調整相對動能']=momentum.get('relative_momentum_risk_adjusted_60',np.nan)
    comparison=ai.get('comparison',pd.DataFrame())
    row['AI同區間最佳模型']=(comparison.iloc[0]['model'] if not comparison.empty else '資料不足')
    for col in ['AI上行先觸機率','AI下行先觸機率','AI盤整機率']: row[col]=np.nan
    row['AI預測期間']=cfg['ai']['horizon']; row['AI上方ATR倍數']=cfg['ai']['up_atr']; row['AI下方ATR倍數']=cfg['ai']['down_atr']
    row['AI目標定義']='未來10日ATR先觸價；同日雙觸採下方'
    for rec in ai['latest'].to_dict('records'):
        if rec['model']=='ensemble':
            col={'up_first':'AI上行先觸機率','down_first':'AI下行先觸機率','neutral':'AI盤整機率'}[rec['class']]
            row[col]=rec['probability']
    if cfg['backtest']['enabled']:
        b,t,e=compare_strategies(raw,adjusted,x,cfg)
    else: b,t,e=pd.DataFrame(),pd.DataFrame(),pd.DataFrame()
    source={'代號':ticker,**meta,'原始列數':len(raw),'還原列數':len(adjusted) if adjusted is not None else 0,
        '動能基準':cfg['data'].get('benchmark','TAIEX'),'動能基準列數':len(benchmark) if benchmark is not None else 0,
        '動能基準資料狀態':momentum.get('reason','資料不足'),
        'TimesFM資料狀態':timesfm_summary['TimesFM狀態'],
        '法人籌碼記錄數':len(chips),'融資融券記錄數':len(chips.attrs.get('margin',[])),
        '借券記錄數':len(chips.attrs.get('lending',[])),'券商分點記錄數':len(chips.attrs.get('branches',[])),
        '最後行情日':str(raw.index[-1].date()),'最後籌碼可用':bool(x.chip_ready.iloc[-1]),
        '回測狀態':'已產出單檔研究比較' if not b.empty else '停用／資料不足'}
    signal=x[['technical_entry','entry','exit','chip_ready','margin_ready','margin_overheat','margin_utilization','lending_ready','lending_high',
              'branch_ready','branch_positive','branch_negative','chip_kline_score','short_trend_up','short_trend_down',
              'medium_trend_up','medium_trend_down','long_trend_up','long_trend_down','trend_up','risk_ok']].copy()
    signal.index.name='date'; signal=signal.reset_index(); signal.insert(0,'代號',ticker)
    return row,source,ai,b,t,e,signal,timesfm_path


def run(config_path='config/timing_TW.yaml', provider=None, candidate_frame=None, provenance=None,
        write=True, progress=print):
    cfg=load_config(config_path)
    project=Path(__file__).resolve().parent.parent
    if candidate_frame is None: candidate_frame,provenance=candidates(cfg,project)
    provider=provider or Provider(cfg)
    provenance=dict(provenance or {'來源':'手動候選名單','WHID評估日期':None})
    provenance.setdefault('第二階段執行日期台北',str(getattr(provider,'report_check_date',provider.asof).date()))
    provenance.setdefault('WHID日期比較規則','以台北日曆日比較WHID報表日期；不與最後交易日比較')
    holdings=read_holdings(cfg['candidates'].get('holdings_path'))
    timesfm_runner=TimesFMPathRunner(cfg,'TW')
    benchmark=None
    if cfg['ai']['enabled']:
        try:benchmark=provider.prices('TaiwanStockPrice',cfg['data'].get('benchmark','TAIEX'))
        except Exception:benchmark=None
    market_summary=market_sentiment(None,cfg)
    if cfg['sentiment']['enabled'] and hasattr(provider,'market_context'):
        try:market_summary=market_sentiment(provider.market_context(),cfg)
        except Exception as exc:market_summary['市場情緒資料狀態']='資料不足：'+type(exc).__name__
    buckets={k:[] for k in ['detail','data_status','ai_latest','ai_metrics','ai_comparison','ai_calibration','ai_predictions',
                            'ai_status','backtest','trades','equity','signals','timesfm_path','news']}
    for i,candidate in enumerate(candidate_frame.to_dict('records'),1):
        ticker=candidate['代號']
        if progress: progress(f'[{i}/{len(candidate_frame)}] {ticker}：資料、規則、AI、TimesFM與回測研究')
        try:
            news_summary,scored_news=news_sentiment(None,ticker,provider.asof,cfg)
            if cfg['sentiment']['enabled'] and hasattr(provider,'news'):
                try:news_summary,scored_news=news_sentiment(provider.news(ticker),ticker,provider.asof,cfg)
                except Exception as exc:news_summary['新聞資料狀態']='資料不足：'+type(exc).__name__
            if not scored_news.empty:buckets['news'].append(scored_news)
            raw,adjusted,chips,meta=provider.bundle(ticker)
            row,source,ai,b,t,e,s,tf_path=analyze_bundle(ticker,candidate,raw,adjusted,chips,meta,cfg,provenance,provider.asof,
                holdings.get(ticker),market_summary,news_summary,benchmark,timesfm_runner)
            buckets['detail'].append(pd.DataFrame([row])); buckets['data_status'].append(pd.DataFrame([source]))
            for key,df in [('ai_latest',ai['latest']),('ai_metrics',ai['metrics']),('ai_comparison',ai['comparison']),('ai_calibration',ai['calibration']),
                ('ai_predictions',ai['predictions']),('backtest',b),('trades',t),('equity',e)]:
                if not df.empty: buckets[key].append(df.assign(代號=ticker))
            buckets['ai_status'].append(pd.DataFrame([{'代號':ticker,'狀態':ai['status'],'說明':'；'.join(ai['errors'])}]))
            buckets['signals'].append(s)
            if not tf_path.empty:buckets['timesfm_path'].append(tf_path)
        except Exception as exc:
            # Do not include raw provider/request exceptions in deliverables.
            reason=type(exc).__name__
            if isinstance(exc,(ValueError,RuntimeError)): reason+=': '+str(exc)[:180]
            buckets['detail'].append(pd.DataFrame([{'代號':ticker,'未持有建議':'資料不足／暫不判斷',
                '持有情境建議':'資料不足，人工確認','AI狀態':'未執行','TimesFM10日路徑判讀':'資料不足',
                'TimesFM狀態':'未執行：股票流程失敗','資料狀態':reason}]))
            buckets['data_status'].append(pd.DataFrame([{'代號':ticker,'錯誤':reason}]))
            if progress: progress(f'{ticker} 未完成：{type(exc).__name__}；已保留失敗列')
    result={k:pd.concat(v,ignore_index=True) if v else pd.DataFrame() for k,v in buckets.items()}
    result['ai_model_summary']=summarize_model_metrics(result['ai_metrics'])
    result['simple']=result['detail'].reindex(columns=SIMPLE_COLUMNS)
    result['market_context']=pd.DataFrame([market_summary])
    result['candidates']=candidate_frame.copy();result['provenance']=provenance
    if write: result['paths']=export_reports(result,cfg)
    return result


def finlab_research_positions(signal_frame):
    """Optional FinLab handoff, not an automatic download/login/trade action.

    Fixed-list historical research only. Weights and stop rules must be supplied
    explicitly in FinLab; this function emits desired long/flat states.
    """
    columns={}
    for ticker,g in signal_frame.groupby('代號'):
        g=g.sort_values('date'); active=False; states=[]
        for rec in g.to_dict('records'):
            if bool(rec['exit']): active=False
            elif bool(rec['entry']): active=True
            states.append(active)
        code=ticker.split('.')[0]
        if code in columns: raise ValueError('重複FinLab代號')
        columns[code]=pd.Series(states,index=pd.to_datetime(g.date))
    return pd.DataFrame(columns).fillna(False).astype(bool)
