"""Independent US stage-2 orchestration; no TW imports, credentials or orders."""
from __future__ import annotations
import json
from datetime import datetime, timezone
from pathlib import Path
import numpy as np
import pandas as pd
from timing_us_data import VERSION, Provider, candidates, load_config, number, read_holdings
from timing_us_rules import signals, latest_assessment, compare_strategies
from timing_us_ai import run_ai

SIMPLE_COLUMNS=['代號','訊號日期NY','分析日收盤價USD','短期趨勢','中期趨勢','長期趨勢','量價參考','市場環境','風險','未持有建議','持有情境建議',
 'ATR風險參考價USD','支撐參考USD','壓力參考USD','AI上行先觸機率','AI下行先觸機率','AI盤整機率','AI狀態','資料狀態']


def safe_json(value):
    if isinstance(value,dict): return {str(k):safe_json(v) for k,v in value.items()}
    if isinstance(value,(list,tuple)): return [safe_json(v) for v in value]
    if isinstance(value,(pd.Timestamp,datetime)): return value.isoformat()
    if isinstance(value,np.integer): return int(value)
    if isinstance(value,np.bool_): return bool(value)
    if isinstance(value,(float,np.floating)): return float(value) if np.isfinite(value) else None
    if value is pd.NA or value is pd.NaT: return None
    return value

def _sheet(frame):
    if frame is None or frame.empty: return pd.DataFrame({'狀態':['本次無資料；請查看DataStatus與AIStatus']})
    f=frame.copy()
    for col in f:
        if f[col].dtype==object or pd.api.types.is_string_dtype(f[col].dtype):
            f[col]=f[col].map(lambda v: json.dumps(safe_json(v),ensure_ascii=False) if isinstance(v,(dict,list)) else v)
            # External strings must never become executable spreadsheet formulas.
            f[col]=f[col].map(lambda v: "'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v)
    return f

def export_reports(result, cfg):
    stamp=datetime.now().strftime('%Y%m%d_%H%M%S_%f')
    run_id=f'US_Timing_{stamp}'
    root=Path(cfg['output']['root'])
    paths={kind:root/kind/datetime.now().strftime('%Y/%m') for kind in ['詳細版','簡化版','研究資料']}
    for p in paths.values(): p.mkdir(parents=True,exist_ok=True)
    simple=paths['簡化版']/(run_id+'_簡化.xlsx'); detailed=paths['詳細版']/(run_id+'_詳細.xlsx')
    rules=[{'分類':g,'設定':k,'值':v} for g in ['rules','ai','backtest'] for k,v in cfg[g].items()]
    rules.extend([
        {'分類':'說明','設定':'趨勢分層','值':'短期5/20/3、中期20/60/5、長期120/240/20；只有中期趨勢參與現行進出場與回測'},
        {'分類':'說明','設定':'AI目標','值':'未來10個交易日先觸及上方1.5倍ATR、下方1.0倍ATR，或期間內兩者皆未觸及；同日雙觸採下方'},
        {'分類':'限制','設定':'AI決策','值':'所有AI機率只作研究參考，未用於交易建議；LSTM為挑戰模型，不加入正式集成'},
        {'分類':'限制','設定':'回測','值':'固定候選名單的單檔研究，不是WHID歷史選股或投資組合績效'},
        {'分類':'限制','設定':'成交','值':'前日訊號下一日開盤；同日雙觸價採停損優先；零量/一價日不成交'},
        {'分類':'限制','設定':'價格','值':'還原價、分數股與股息再投資研究假設；未模擬實際張數/最低手續費'},
        {'分類':'限制','設定':'量價參考','值':'CMF與帶方向成交量是價量代理指標，並非機構資金流；沒有使用13F季度持股'},
        {'分類':'限制','設定':'成本','值':'USD每邊佣金預留、賣出額外費用及滑價為估計；不是券商實價或法定稅率，不計個人稅負'},
        {'分類':'限制','設定':'資訊時間','值':'紐約時區依設定截止小時排除未完整日線，使用常規交易時段；處理夏令時間'},
        {'分類':'限制','設定':'近期參考','值':'最新日缺值可顯示近期完整日的歷史參考；日期及價格明確標示，未產生最新交易建議。可在data設定停用。'},
        {'分類':'限制','設定':'風險參考價','值':'目前收盤價減ATR倍數；不是保證停損成交價，也不是個人化部位建議'},
    ])
    sheets={'Report':result['detail'],'DataStatus':result['data_status'],'Backtest':result['backtest'],
        'Trades':result['trades'],'Equity':result['equity'],'AILatest':result['ai_latest'],
        'AIMetrics':result['ai_metrics'],'AICalibration':result['ai_calibration'],'AIStatus':result['ai_status'],
        'Rules':pd.DataFrame(rules),'CandidateSource':pd.DataFrame([result['provenance']]),
        'Candidates':result['candidates']}
    # This is the program's reusable export function, not a change to any existing workbook.
    from openpyxl.styles import Font,PatternFill,Alignment
    for path,book in [(simple,{'Report':result['simple']}),(detailed,sheets)]:
        with pd.ExcelWriter(path,engine='openpyxl') as writer:
            for name,df in book.items():
                out=_sheet(df)
                out.to_excel(writer,sheet_name=name,index=False)
                ws=writer.sheets[name]; ws.freeze_panes='C2'; ws.auto_filter.ref=ws.dimensions
                for cell in ws[1]:
                    cell.font=Font(color='FFFFFF',bold=True); cell.fill=PatternFill('solid',fgColor='17365D')
                    cell.alignment=Alignment(wrap_text=True,vertical='center')
                ws.row_dimensions[1].height=32
                for j,(col,cells) in enumerate(zip(out.columns,ws.iter_cols()),1):
                    from openpyxl.utils import get_column_letter
                    ws.column_dimensions[get_column_letter(j)].width=min(42,max(14,len(str(col))*1.7))
                    for cell in list(cells)[1:]:
                        if isinstance(cell.value,(int,float)):
                            cell.number_format='0.0%' if any(s in str(col).lower() for s in ['probability','return','drawdown','win_rate','比例','機率','報酬率','乖離']) else '#,##0.000'
    snapshot=paths['研究資料']/(run_id+'_snapshot.json')
    snapshot.write_text(json.dumps(safe_json({'version':VERSION,'created_utc':datetime.now(timezone.utc),
        'config':cfg,'provenance':result['provenance'],'rows':result['detail'].to_dict('records')}),ensure_ascii=False,indent=2),encoding='utf-8')
    # Preserve per-day out-of-sample outputs without bloating the simplified report.
    if not result['ai_predictions'].empty:
        result['ai_predictions'].to_csv(paths['研究資料']/(run_id+'_ai_oos.csv'),index=False,encoding='utf-8-sig')
    result['signals'].to_csv(paths['研究資料']/(run_id+'_signals.csv'),index=False,encoding='utf-8-sig')
    return {'簡化版':str(simple),'詳細版':str(detailed),'快照':str(snapshot)}

def select_analysis_window(raw, adjusted, benchmark, cfg, asof):
    """Return an explicitly dated reference window, never a fabricated latest bar."""
    status={'要求行情日期NY':str(raw.index[-1].date()),'行情分析模式':'最新日分析','參考間隔交易日期數':0}
    cols=['Open','High','Low','Close','Volume']
    valid=np.isfinite(raw[cols]).all(axis=1)&(raw.Close>0)&(raw.Volume>0)
    if adjusted is None:valid[:]=False
    else:valid&=np.isfinite(adjusted[cols]).all(axis=1)&(adjusted.Close>0)
    if valid.iloc[-1]:return raw,adjusted,benchmark,status
    status['行情分析模式']='最新日資料不足'
    if not cfg['data'].get('allow_recent_reference',True):return raw,adjusted,benchmark,status
    if cfg['rules']['require_market_for_entry']:
        if benchmark is None:return raw,adjusted,benchmark,status
        b=benchmark.reindex(raw.index)
        valid&=np.isfinite(b[cols]).all(axis=1)&(b.Close>0)&(b.Volume>0)
    positions=np.flatnonzero(valid.to_numpy())
    if len(positions)==0:return raw,adjusted,benchmark,status
    end=int(positions[-1]);day=raw.index[end];gap=len(raw)-end-1
    age=(pd.Timestamp(asof).normalize()-day).days
    if gap>cfg['data'].get('max_reference_gap_sessions',2) or not 0<=age<=cfg['data']['max_price_age_days']:
        return raw,adjusted,benchmark,status
    status.update({'行情分析模式':'最近完整日歷史參考','參考間隔交易日期數':gap})
    return raw.iloc[:end+1],adjusted.iloc[:end+1],benchmark.loc[:day] if benchmark is not None else None,status


def analyze_bundle(ticker, candidate, raw, adjusted, benchmark, meta, cfg, provenance, asof, holding=None):
    raw,adjusted,benchmark,window=select_analysis_window(raw,adjusted,benchmark,cfg,asof)
    x=signals(raw,adjusted,benchmark,cfg)
    row=latest_assessment(ticker,raw,adjusted,x,cfg,provenance,asof,holding)
    row['分析日收盤價USD']=row.pop('現價USD')
    for col,value in candidate.items():
        if col!='代號': row['WHID_'+col]=value
    row.update(meta);row.update(window)
    reference=window['行情分析模式']=='最近完整日歷史參考'
    if reference:
        row['參考日未持有判讀']=row['未持有建議'];row['參考日持有判讀']=row['持有情境建議']
        row['未持有建議']='歷史參考：'+row['未持有建議']+'；待最新行情確認'
        row['持有情境建議']='歷史參考：'+row['持有情境建議']+'；待最新行情確認'
        row['資料狀態']=f"截至{raw.index[-1].date()}歷史參考；{window['要求行情日期NY']}行情不完整，未產生最新交易建議；"+row['資料狀態']
    stale=(pd.Timestamp(asof).normalize()-raw.index[-1]).days>cfg['data']['max_price_age_days']
    if stale:
        ai={'latest':pd.DataFrame(),'metrics':pd.DataFrame(),'calibration':pd.DataFrame(),
            'predictions':pd.DataFrame(),'errors':[],'status':'行情過期，AI未執行'}
    else: ai=run_ai(adjusted,x,cfg)
    row['AI狀態']=('參考日模型，非最新預測；' if reference else '')+ai['status']
    for col in ['AI上行先觸機率','AI下行先觸機率','AI盤整機率']: row[col]=np.nan
    row['AI預測期間']=cfg['ai']['horizon']; row['AI上方ATR倍數']=cfg['ai']['up_atr']; row['AI下方ATR倍數']=cfg['ai']['down_atr']
    row['AI目標定義']='未來10日ATR先觸價；同日雙觸採下方'
    for rec in ai['latest'].to_dict('records'):
        if rec['model']=='ensemble':
            col={'up_first':'AI上行先觸機率','down_first':'AI下行先觸機率','neutral':'AI盤整機率'}[rec['class']]
            row[col]=rec['probability']
    if cfg['backtest']['enabled']:
        b,t,e=compare_strategies(raw,adjusted,x,cfg,benchmark)
    else: b,t,e=pd.DataFrame(),pd.DataFrame(),pd.DataFrame()
    source={'代號':ticker,**meta,**window,'原始列數':len(raw),'還原列數':len(adjusted) if adjusted is not None else 0,
        '基準列數':len(benchmark) if benchmark is not None else 0,'最後行情日':str(raw.index[-1].date()),'量價指標可用':bool(x.flow_ready.iloc[-1]),
        '回測狀態':'已產出單檔研究比較' if not b.empty else '停用／資料不足'}
    signal=x[['technical_entry','entry','exit','flow_ready','market_ready','short_trend_up','short_trend_down',
              'medium_trend_up','medium_trend_down','long_trend_up','long_trend_down','trend_up','risk_ok']].copy()
    signal.index.name='date'; signal=signal.reset_index(); signal.insert(0,'代號',ticker)
    return row,source,ai,b,t,e,signal

def run(config_path='config/timing_US.yaml', provider=None, candidate_frame=None, provenance=None,
        write=True, progress=print):
    cfg=load_config(config_path)
    project=Path(__file__).resolve().parent.parent
    if candidate_frame is None: candidate_frame,provenance=candidates(cfg,project)
    provider=provider or Provider(cfg)
    provenance=dict(provenance or {'來源':'手動候選名單','WHID評估日期':None})
    provenance.setdefault('第二階段執行日期台北',str(getattr(provider,'report_check_date',provider.asof).date()))
    provenance.setdefault('WHID日期比較規則','以台北日曆日比較WHID報表日期；不與最後交易日比較')
    holdings=read_holdings(cfg['candidates'].get('holdings_path'))
    buckets={k:[] for k in ['detail','data_status','ai_latest','ai_metrics','ai_calibration','ai_predictions',
                            'ai_status','backtest','trades','equity','signals']}
    for i,candidate in enumerate(candidate_frame.to_dict('records'),1):
        ticker=candidate['代號']
        if progress: progress(f'[{i}/{len(candidate_frame)}] {ticker}：資料、規則、AI研究與回測')
        try:
            raw,adjusted,benchmark,meta=provider.bundle(ticker)
            row,source,ai,b,t,e,s=analyze_bundle(ticker,candidate,raw,adjusted,benchmark,meta,cfg,provenance,provider.asof,holdings.get(ticker))
            buckets['detail'].append(pd.DataFrame([row])); buckets['data_status'].append(pd.DataFrame([source]))
            for key,df in [('ai_latest',ai['latest']),('ai_metrics',ai['metrics']),('ai_calibration',ai['calibration']),
                ('ai_predictions',ai['predictions']),('backtest',b),('trades',t),('equity',e)]:
                if not df.empty: buckets[key].append(df.assign(代號=ticker))
            buckets['ai_status'].append(pd.DataFrame([{'代號':ticker,'狀態':ai['status'],'說明':'；'.join(ai['errors'])}]))
            buckets['signals'].append(s)
        except Exception as exc:
            # Do not include raw provider/request exceptions in deliverables.
            reason=type(exc).__name__
            if isinstance(exc,(ValueError,RuntimeError)): reason+=': '+str(exc)[:180]
            buckets['detail'].append(pd.DataFrame([{'代號':ticker,'未持有建議':'資料不足／暫不判斷',
                '持有情境建議':'資料不足，人工確認','AI狀態':'未執行','資料狀態':reason}]))
            buckets['data_status'].append(pd.DataFrame([{'代號':ticker,'錯誤':reason}]))
            if progress: progress(f'{ticker} 未完成：{type(exc).__name__}；已保留失敗列')
    result={k:pd.concat(v,ignore_index=True) if v else pd.DataFrame() for k,v in buckets.items()}
    result['simple']=result['detail'].reindex(columns=SIMPLE_COLUMNS)
    result['candidates']=candidate_frame.copy();result['provenance']=provenance
    if write: result['paths']=export_reports(result,cfg)
    return result

def finlab_research_positions(signal_frame):
    """Optional FinLab handoff, not an automatic download/login/trade action.

    Fixed-list historical research only. Weights and stop rules must be supplied
    explicitly in FinLab; this function emits desired long/flat states.
    """
    columns={}
    for ticker,g in signal_frame.groupby('代號'):
        g=g.sort_values('date'); active=False; states=[]
        for rec in g.to_dict('records'):
            if bool(rec['exit']): active=False
            elif bool(rec['entry']): active=True
            states.append(active)
        code=ticker
        if code in columns: raise ValueError('重複FinLab代號')
        columns[code]=pd.Series(states,index=pd.to_datetime(g.date))
    return pd.DataFrame(columns).fillna(False).astype(bool)
