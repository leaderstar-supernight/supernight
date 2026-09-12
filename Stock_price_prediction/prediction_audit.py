"""Immutable daily forecast ledger, delayed outcome settlement, and 30-day review."""
from __future__ import annotations

import argparse
import json
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd

BASE = Path(__file__).resolve().parent
PROJECT = BASE.parent
VALIDATION_ROOT = PROJECT / 'system_data' / 'validation'
DB_PATH = VALIDATION_ROOT / 'prediction_audit.sqlite3'
REVIEW_ROOT = PROJECT / 'reports' / '30日評估'
CLASS_ORDER = ['down_first', 'neutral', 'up_first']
TIMESFM_CLASS_MAP = {
    '下行先觸': 'down_first',
    '盤整': 'neutral',
    '上行先觸': 'up_first',
    'down_first': 'down_first',
    'neutral': 'neutral',
    'up_first': 'up_first',
}


def _json_value(value):
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.isoformat()
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating, float)):
        return float(value) if np.isfinite(value) else None
    if isinstance(value, np.bool_):
        return bool(value)
    if value is pd.NA or value is pd.NaT:
        return None
    return value


def connect(db_path=DB_PATH):
    db_path = Path(db_path)
    db_path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(db_path)
    connection.execute('PRAGMA journal_mode=WAL')
    connection.executescript('''
    CREATE TABLE IF NOT EXISTS forecasts (
        market TEXT NOT NULL,
        ticker TEXT NOT NULL,
        signal_date TEXT NOT NULL,
        model TEXT NOT NULL,
        recorded_utc TEXT NOT NULL,
        source_report TEXT NOT NULL,
        horizon INTEGER NOT NULL,
        p_down REAL NOT NULL,
        p_neutral REAL NOT NULL,
        p_up REAL NOT NULL,
        signal_close REAL NOT NULL,
        signal_atr REAL NOT NULL,
        upper_barrier REAL NOT NULL,
        lower_barrier REAL NOT NULL,
        actual_class TEXT,
        maturity_date TEXT,
        settled_utc TEXT,
        settlement_status TEXT NOT NULL DEFAULT 'pending',
        PRIMARY KEY (market, ticker, signal_date, model)
    );
    CREATE TABLE IF NOT EXISTS indicators (
        market TEXT NOT NULL,
        ticker TEXT NOT NULL,
        signal_date TEXT NOT NULL,
        recorded_utc TEXT NOT NULL,
        source_report TEXT NOT NULL,
        snapshot_json TEXT NOT NULL,
        horizon INTEGER NOT NULL,
        signal_close REAL,
        signal_atr REAL,
        upper_barrier REAL,
        lower_barrier REAL,
        actual_class TEXT,
        maturity_date TEXT,
        return_5 REAL,
        return_10 REAL,
        return_20 REAL,
        benchmark_return_5 REAL,
        benchmark_return_10 REAL,
        benchmark_return_20 REAL,
        abnormal_return_5 REAL,
        abnormal_return_10 REAL,
        abnormal_return_20 REAL,
        return_60 REAL,
        return_126 REAL,
        return_252 REAL,
        max_up_10 REAL,
        max_down_10 REAL,
        settled_utc TEXT,
        settlement_status TEXT NOT NULL DEFAULT 'pending',
        PRIMARY KEY (market, ticker, signal_date)
    );
    CREATE TABLE IF NOT EXISTS audit_runs (
        run_utc TEXT PRIMARY KEY,
        action TEXT NOT NULL,
        status TEXT NOT NULL,
        details TEXT
    );
    ''')
    existing = {row[1] for row in connection.execute('PRAGMA table_info(indicators)')}
    for column in ['return_60', 'return_126', 'return_252',
                   'benchmark_return_5','benchmark_return_10','benchmark_return_20',
                   'abnormal_return_5','abnormal_return_10','abnormal_return_20']:
        if column not in existing:
            connection.execute(f'ALTER TABLE indicators ADD COLUMN {column} REAL')
    connection.commit()
    return connection


def latest_detailed_report(market, not_before=None):
    folder = PROJECT / 'reports' / '詳細版' / market
    files = [p for p in folder.glob(f'{market}_*_Step2_Report.xlsx')
             if not p.name.startswith('~$')] if folder.exists() else []
    if not_before is not None:
        threshold = pd.Timestamp(not_before).timestamp()
        files = [p for p in files if p.stat().st_mtime >= threshold]
    return max(files, key=lambda p: p.stat().st_mtime) if files else None


def _number(value):
    try:
        result = float(value)
        return result if np.isfinite(result) else None
    except (TypeError, ValueError):
        return None


def _first(row, names):
    for name in names:
        if name in row and pd.notna(row[name]):
            return row[name]
    return None


def archive_report(market, report_path, db_path=DB_PATH):
    """Archive the first forecast for each market/ticker/signal date; never overwrite it."""
    report_path = Path(report_path).resolve()
    with pd.ExcelFile(report_path) as book:
        report = pd.read_excel(book, sheet_name='Report')
        latest = pd.read_excel(book, sheet_name='AILatest')
    required = {'代號', 'prediction_date', 'model', 'class', 'probability', 'signal_close',
                'signal_atr', 'upper_barrier', 'lower_barrier', 'horizon'}
    missing = required - set(latest.columns)
    warnings = []
    if missing:
        warnings.append(f'AILatest無可封存的新版三分類結果，缺少：{sorted(missing)}')
        latest = pd.DataFrame(columns=sorted(required))
    recorded = datetime.now(timezone.utc).isoformat()
    inserted_forecasts = 0
    inserted_indicators = 0
    connection = connect(db_path)
    try:
        for (ticker, signal_date, model), group in latest.groupby(['代號', 'prediction_date', 'model'], dropna=False):
            probability = group.set_index('class').probability.reindex(CLASS_ORDER)
            if probability.isna().any() or abs(float(probability.sum()) - 1.0) > 1e-5:
                continue
            first = group.iloc[0]
            cursor = connection.execute('''
                INSERT OR IGNORE INTO forecasts
                (market,ticker,signal_date,model,recorded_utc,source_report,horizon,
                 p_down,p_neutral,p_up,signal_close,signal_atr,upper_barrier,lower_barrier)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)''', (
                market, str(ticker), str(pd.Timestamp(signal_date).date()), str(model), recorded,
                str(report_path), int(first.horizon), float(probability.down_first),
                float(probability.neutral), float(probability.up_first), float(first.signal_close),
                float(first.signal_atr), float(first.upper_barrier), float(first.lower_barrier)))
            inserted_forecasts += cursor.rowcount
        report_date_col = '訊號日期' if market == 'TW' else '訊號日期NY'
        close_col = '現價' if market == 'TW' else '分析日收盤價USD'
        ensemble_source = latest.loc[latest.model == 'ensemble'].copy()
        ensemble_source['prediction_date'] = pd.to_datetime(ensemble_source.prediction_date).dt.strftime('%Y-%m-%d')
        ensemble = ensemble_source.groupby(['代號', 'prediction_date']).first()
        for _, row in report.iterrows():
            ticker = str(row.get('代號', '')).strip()
            if not ticker or pd.isna(row.get(report_date_col)):
                continue
            signal_date = str(pd.Timestamp(row[report_date_col]).date())
            key = (ticker, signal_date)
            ai_row = ensemble.loc[key] if key in ensemble.index else None
            close = _number(ai_row.signal_close) if ai_row is not None else _number(row.get(close_col))
            atr = _number(ai_row.signal_atr) if ai_row is not None else None
            if atr is None and close is not None:
                atr_pct = _number(row.get('ATR比例'))
                atr = close * atr_pct if atr_pct is not None else None
            upper = _number(ai_row.upper_barrier) if ai_row is not None else (close + 1.5 * atr if close and atr else None)
            lower = _number(ai_row.lower_barrier) if ai_row is not None else (close - atr if close and atr else None)
            snapshot = {str(k): _json_value(v) for k, v in row.to_dict().items()}
            cursor = connection.execute('''
                INSERT OR IGNORE INTO indicators
                (market,ticker,signal_date,recorded_utc,source_report,snapshot_json,horizon,
                 signal_close,signal_atr,upper_barrier,lower_barrier)
                VALUES (?,?,?,?,?,?,?,?,?,?,?)''', (
                market, ticker, signal_date, recorded, str(report_path),
                json.dumps(snapshot, ensure_ascii=False), 10, close, atr, upper, lower))
            inserted_indicators += cursor.rowcount
        connection.execute('INSERT INTO audit_runs VALUES (?,?,?,?)', (
            recorded, f'archive_{market}', 'ok', json.dumps({
                'report': str(report_path), 'forecasts': inserted_forecasts,
                'indicators': inserted_indicators, 'warnings': warnings}, ensure_ascii=False)))
        connection.commit()
    finally:
        connection.close()
    return {'market': market, 'report': str(report_path), 'forecasts': inserted_forecasts,
            'indicators': inserted_indicators, 'warnings': warnings}


def archive_latest(market, not_before=None, db_path=DB_PATH):
    path = latest_detailed_report(market, not_before)
    if path is None:
        raise FileNotFoundError(f'找不到{market}本次新產生的第二階段詳細報表')
    return archive_report(market, path, db_path)


def _provider_prices(market, ticker):
    if market == 'TW':
        from timing_data import Provider, load_config
        cfg = load_config(BASE / 'config' / 'timing_TW.yaml')
        provider=Provider(cfg)
        raw, adjusted, _, _ = provider.bundle(ticker)
        benchmark=provider.prices('TaiwanStockPrice','TAIEX')
    else:
        from timing_us_data import Provider, load_config
        cfg = load_config(BASE / 'config' / 'timing_US.yaml')
        raw, adjusted, benchmark, _ = Provider(cfg).bundle(ticker)
    if raw is None or adjusted is None:
        raise RuntimeError('原始價或還原價不可用')
    return raw.sort_index(), adjusted.sort_index(), benchmark.sort_index() if benchmark is not None else None


def _outcome(raw, adjusted, signal_date, horizon, upper, lower, benchmark=None):
    signal_date = pd.Timestamp(signal_date)
    future_raw = raw.loc[raw.index > signal_date].head(max(252, horizon))
    if len(future_raw) < horizon:
        return None
    signal_adjusted = adjusted.loc[adjusted.index <= signal_date]
    future_adjusted = adjusted.loc[adjusted.index > signal_date].head(252)
    if signal_adjusted.empty or len(future_adjusted) < horizon:
        return None
    actual = 'neutral'
    for _, bar in future_raw.head(horizon).iterrows():
        if np.isfinite(bar.Low) and bar.Low <= lower:
            actual = 'down_first'
            break
        if np.isfinite(bar.High) and bar.High >= upper:
            actual = 'up_first'
            break
    base = float(signal_adjusted.Close.iloc[-1])
    def forward_return(days):
        return (float(future_adjusted.Close.iloc[days - 1]) / base - 1) if len(future_adjusted) >= days else None
    benchmark_base=None
    if benchmark is not None:
        prior=benchmark.loc[benchmark.index<=signal_date,'Close'].dropna()
        benchmark_base=float(prior.iloc[-1]) if len(prior) else None
    def benchmark_return(days):
        if benchmark_base is None or len(future_adjusted)<days:return None
        target=future_adjusted.index[days-1]
        observed=benchmark.loc[(benchmark.index>signal_date)&(benchmark.index<=target),'Close'].dropna()
        return float(observed.iloc[-1])/benchmark_base-1 if len(observed) else None
    stock_returns={d:forward_return(d) for d in [5,10,20,60,126,252]}
    benchmark_returns={d:benchmark_return(d) for d in [5,10,20]}
    signal_raw = raw.loc[raw.index <= signal_date]
    raw_base = float(signal_raw.Close.iloc[-1]) if not signal_raw.empty else np.nan
    ten = future_raw.head(horizon)
    return {
        'actual_class': actual, 'maturity_date': str(ten.index[-1].date()),
        'return_5':stock_returns[5], 'return_10':stock_returns[10], 'return_20':stock_returns[20],
        'benchmark_return_5':benchmark_returns[5], 'benchmark_return_10':benchmark_returns[10],
        'benchmark_return_20':benchmark_returns[20],
        'abnormal_return_5':stock_returns[5]-benchmark_returns[5] if stock_returns[5] is not None and benchmark_returns[5] is not None else None,
        'abnormal_return_10':stock_returns[10]-benchmark_returns[10] if stock_returns[10] is not None and benchmark_returns[10] is not None else None,
        'abnormal_return_20':stock_returns[20]-benchmark_returns[20] if stock_returns[20] is not None and benchmark_returns[20] is not None else None,
        'return_60':stock_returns[60], 'return_126':stock_returns[126], 'return_252':stock_returns[252],
        'max_up_10': float(ten.High.max() / raw_base - 1) if np.isfinite(raw_base) else None,
        'max_down_10': float(ten.Low.min() / raw_base - 1) if np.isfinite(raw_base) else None,
    }


def settle_pending(db_path=DB_PATH):
    connection = connect(db_path)
    forecasts = pd.read_sql_query(
        "SELECT * FROM forecasts WHERE settlement_status='pending'", connection)
    indicators = pd.read_sql_query(
        "SELECT * FROM indicators WHERE settlement_status='pending' OR return_252 IS NULL OR abnormal_return_10 IS NULL", connection)
    keys = sorted(set(zip(forecasts.market, forecasts.ticker)) |
                  set(zip(indicators.market, indicators.ticker)))
    settled_forecasts = settled_indicators = enriched_indicators = 0
    errors = []
    now = datetime.now(timezone.utc).isoformat()
    for market, ticker in keys:
        try:
            raw, adjusted, benchmark = _provider_prices(market, ticker)
        except Exception as exc:
            errors.append(f'{market}/{ticker}: {type(exc).__name__}: {str(exc)[:140]}')
            continue
        subset = forecasts.loc[(forecasts.market == market) & (forecasts.ticker == ticker)]
        for _, row in subset.iterrows():
            result = _outcome(raw, adjusted, row.signal_date, int(row.horizon),
                              float(row.upper_barrier), float(row.lower_barrier), benchmark)
            if result is None:
                continue
            connection.execute('''UPDATE forecasts SET actual_class=?,maturity_date=?,settled_utc=?,
                settlement_status='settled' WHERE market=? AND ticker=? AND signal_date=? AND model=?''',
                (result['actual_class'], result['maturity_date'], now, market, ticker,
                 row.signal_date, row.model))
            settled_forecasts += 1
        subset = indicators.loc[(indicators.market == market) & (indicators.ticker == ticker)]
        for _, row in subset.iterrows():
            if not all(np.isfinite(_number(v) or np.nan) for v in [row.upper_barrier, row.lower_barrier]):
                connection.execute('''UPDATE indicators SET settlement_status='missing_barrier'
                    WHERE market=? AND ticker=? AND signal_date=?''', (market, ticker, row.signal_date))
                continue
            result = _outcome(raw, adjusted, row.signal_date, int(row.horizon),
                              float(row.upper_barrier), float(row.lower_barrier), benchmark)
            if result is None:
                continue
            connection.execute('''UPDATE indicators SET actual_class=?,maturity_date=?,return_5=?,return_10=?,
                return_20=?,benchmark_return_5=?,benchmark_return_10=?,benchmark_return_20=?,
                abnormal_return_5=?,abnormal_return_10=?,abnormal_return_20=?,
                return_60=?,return_126=?,return_252=?,max_up_10=?,max_down_10=?,
                settled_utc=?,settlement_status='settled'
                WHERE market=? AND ticker=? AND signal_date=?''',
                (result['actual_class'], result['maturity_date'], result['return_5'], result['return_10'],
                 result['return_20'],result['benchmark_return_5'],result['benchmark_return_10'],result['benchmark_return_20'],
                 result['abnormal_return_5'],result['abnormal_return_10'],result['abnormal_return_20'],
                 result['return_60'], result['return_126'], result['return_252'],
                 result['max_up_10'], result['max_down_10'], now,
                 market, ticker, row.signal_date))
            if row.settlement_status == 'pending':
                settled_indicators += 1
            elif any(pd.isna(row.get(column)) and result[column] is not None
                     for column in ['return_20','return_60','return_126','return_252',
                                    'benchmark_return_10','abnormal_return_10']):
                enriched_indicators += 1
    connection.execute('INSERT INTO audit_runs VALUES (?,?,?,?)', (
        now, 'settle', 'ok' if not errors else 'partial', json.dumps({
            'forecasts': settled_forecasts, 'indicators': settled_indicators,
            'long_horizon_updates': enriched_indicators,
            'errors': errors}, ensure_ascii=False)))
    connection.commit()
    connection.close()
    return {'forecasts': settled_forecasts, 'indicators': settled_indicators,
            'long_horizon_updates': enriched_indicators, 'errors': errors}


def _ai_tables(frame):
    if frame.empty:
        return pd.DataFrame(), pd.DataFrame()
    from sklearn.metrics import accuracy_score, balanced_accuracy_score, f1_score, log_loss
    rows, calibration = [], []
    label_map = {name: i for i, name in enumerate(CLASS_ORDER)}
    for (market, model), group in frame.groupby(['market', 'model']):
        y = group.actual_class.map(label_map).to_numpy(dtype=int)
        probability = group[['p_down', 'p_neutral', 'p_up']].to_numpy(dtype=float)
        probability = np.clip(probability, 1e-9, 1)
        probability /= probability.sum(axis=1, keepdims=True)
        pred = probability.argmax(axis=1)
        one_hot = np.eye(3)[y]
        base_probability = np.bincount(y, minlength=3) / len(y)
        base = np.tile(base_probability, (len(y), 1))
        brier = float(np.mean(np.sum((probability - one_hot) ** 2, axis=1)))
        base_brier = float(np.mean(np.sum((base - one_hot) ** 2, axis=1)))
        marginal_error = float(np.mean(np.abs(probability.mean(axis=0) - one_hot.mean(axis=0))))
        macro_f1 = float(f1_score(y, pred, labels=[0, 1, 2], average='macro', zero_division=0))
        verdict = '資料不足'
        if len(y) >= 30:
            skill = 1 - brier / base_brier if base_brier else np.nan
            if skill > 0 and marginal_error <= .10:
                verdict = '初步有參考性'
            elif skill > 0:
                verdict = '有區辨力；校準需修正'
            else:
                verdict = '未優於基準；需修正'
        rows.append({
            '市場': market, '模型': model, '成熟樣本': len(y),
            '正確率': float(accuracy_score(y, pred)),
            '平衡正確率': float(balanced_accuracy_score(y, pred)) if len(np.unique(y)) > 1 else np.nan,
            'Macro F1': macro_f1, 'Log Loss': float(log_loss(y, probability, labels=[0, 1, 2])),
            '基準Log Loss': float(log_loss(y, base, labels=[0, 1, 2])),
            'Brier': brier, '基準Brier': base_brier,
            'Brier Skill': 1 - brier / base_brier if base_brier else np.nan,
            '邊際校準誤差': marginal_error, '判定': verdict,
        })
        for class_id, class_name in enumerate(CLASS_ORDER):
            for lower in np.arange(0, 1, .1):
                upper = lower + .1
                selected = (probability[:, class_id] >= lower) & ((probability[:, class_id] < upper) if upper < 1 else True)
                if not selected.any():
                    continue
                calibration.append({
                    '市場': market, '模型': model, '分類': class_name,
                    '機率區間': f'{lower:.0%}–{upper:.0%}', '樣本數': int(selected.sum()),
                    '平均預測機率': float(probability[selected, class_id].mean()),
                    '實際發生率': float((y[selected] == class_id).mean()),
                })
    return pd.DataFrame(rows), pd.DataFrame(calibration)


def _timesfm_table(frame):
    """Evaluate TimesFM point paths and quantile-derived class probabilities."""
    columns = [
        '市場', '模型', '成熟樣本', '點路徑正確率', '機率成熟樣本',
        '機率最高類正確率', '機率平衡正確率', '機率Macro F1',
        '多數類基準正確率', '相對基準差',
        '平衡正確率', 'Macro F1', '下行先觸召回率', '盤整召回率',
        '上行先觸召回率', '10日預測報酬MAE', '10日預測報酬RMSE',
        '10日預測報酬偏誤', '10日方向命中率', '預測與實際報酬Spearman',
        'Log Loss', '基準Log Loss', 'Brier', '基準Brier', 'Brier Skill',
        '邊際校準誤差', '判定',
    ]
    if frame.empty:
        return pd.DataFrame(columns=columns)
    from scipy.stats import spearmanr
    from sklearn.metrics import (accuracy_score, balanced_accuracy_score, f1_score,
                                 log_loss, recall_score)

    records = []
    for _, outcome in frame.reset_index(drop=True).iterrows():
        try:
            snapshot = json.loads(outcome.snapshot_json)
        except (TypeError, ValueError, json.JSONDecodeError):
            continue
        predicted_class = TIMESFM_CLASS_MAP.get(str(snapshot.get('TimesFM10日路徑判讀', '')).strip())
        actual_class = str(outcome.get('actual_class', '')).strip()
        status_text = str(snapshot.get('TimesFM狀態', ''))
        if predicted_class is None or actual_class not in CLASS_ORDER or '可用' not in status_text:
            continue
        records.append({
            'market': outcome.market,
            'model': str(snapshot.get('TimesFM模型') or 'TimesFM'),
            'predicted_class': predicted_class,
            'actual_class': actual_class,
            'predicted_return': _number(snapshot.get('TimesFM10日預測漲跌幅')),
            'actual_return': _number(outcome.get('return_10')),
            'p_down': _number(snapshot.get('TimesFM估計下行先觸機率')),
            'p_neutral': _number(snapshot.get('TimesFM估計盤整機率')),
            'p_up': _number(snapshot.get('TimesFM估計上行先觸機率')),
        })
    samples = pd.DataFrame(records)
    if samples.empty:
        return pd.DataFrame(columns=columns)

    rows = []
    label_map = {name: i for i, name in enumerate(CLASS_ORDER)}
    for (market, model), group in samples.groupby(['market', 'model'], dropna=False):
        y = group.actual_class.map(label_map).to_numpy(dtype=int)
        pred = group.predicted_class.map(label_map).to_numpy(dtype=int)
        point_accuracy = float(accuracy_score(y, pred))
        majority_accuracy = float(pd.Series(y).value_counts(normalize=True).max())
        recalls = recall_score(y, pred, labels=[0, 1, 2], average=None, zero_division=0)
        point = group.dropna(subset=['predicted_return', 'actual_return']).copy()
        if point.empty:
            mae = rmse = bias = direction = rho = np.nan
        else:
            error = point.predicted_return - point.actual_return
            mae = float(error.abs().mean())
            rmse = float(np.sqrt(np.mean(np.square(error))))
            bias = float(error.mean())
            direction = float((np.sign(point.predicted_return) == np.sign(point.actual_return)).mean())
            rho = np.nan
            if len(point) >= 3 and point.predicted_return.nunique() > 1 and point.actual_return.nunique() > 1:
                rho = float(spearmanr(point.predicted_return, point.actual_return).statistic)
        balanced = (float(balanced_accuracy_score(y, pred))
                    if len(np.unique(y)) > 1 else np.nan)
        probabilistic = group.dropna(subset=['p_down', 'p_neutral', 'p_up']).copy()
        probabilistic = probabilistic.loc[
            (probabilistic[['p_down', 'p_neutral', 'p_up']] >= 0).all(axis=1)
            & (probabilistic[['p_down', 'p_neutral', 'p_up']].sum(axis=1) > 0)
        ]
        probability_accuracy = probability_balanced = probability_f1 = np.nan
        logloss = base_logloss = brier = base_brier = skill = marginal_error = np.nan
        if not probabilistic.empty:
            y_probability = probabilistic.actual_class.map(label_map).to_numpy(dtype=int)
            probability = probabilistic[['p_down', 'p_neutral', 'p_up']].to_numpy(dtype=float)
            probability = np.clip(probability, 1e-9, 1)
            probability /= probability.sum(axis=1, keepdims=True)
            probability_pred = probability.argmax(axis=1)
            one_hot = np.eye(3)[y_probability]
            base_probability = np.bincount(y_probability, minlength=3) / len(y_probability)
            base = np.tile(base_probability, (len(y_probability), 1))
            probability_accuracy = float(accuracy_score(y_probability, probability_pred))
            probability_balanced = (float(balanced_accuracy_score(y_probability, probability_pred))
                                    if len(np.unique(y_probability)) > 1 else np.nan)
            probability_f1 = float(f1_score(
                y_probability, probability_pred, labels=[0, 1, 2],
                average='macro', zero_division=0))
            logloss = float(log_loss(y_probability, probability, labels=[0, 1, 2]))
            base_logloss = float(log_loss(y_probability, base, labels=[0, 1, 2]))
            brier = float(np.mean(np.sum((probability - one_hot) ** 2, axis=1)))
            base_brier = float(np.mean(np.sum((base - one_hot) ** 2, axis=1)))
            skill = 1 - brier / base_brier if base_brier else np.nan
            marginal_error = float(np.mean(np.abs(probability.mean(axis=0) - one_hot.mean(axis=0))))
        verdict = '資料不足'
        if len(group) >= 30:
            if len(probabilistic) >= 30 and np.isfinite(skill) and skill > 0 and marginal_error <= .10:
                verdict = '估計機率初步有參考性'
            elif len(probabilistic) >= 30 and np.isfinite(skill) and skill > 0:
                verdict = '有區辨力；估計機率校準需修正'
            elif np.isfinite(balanced) and balanced >= .40 and point_accuracy > majority_accuracy:
                verdict = '點路徑初步有參考性；機率待累積'
            elif point_accuracy > majority_accuracy:
                verdict = '點路徑有區辨力；類別平衡需觀察'
            else:
                verdict = '未優於多數類基準；需修正'
        rows.append({
            '市場': market, '模型': model, '成熟樣本': len(group),
            '點路徑正確率': point_accuracy, '機率成熟樣本': len(probabilistic),
            '機率最高類正確率': probability_accuracy,
            '機率平衡正確率': probability_balanced, '機率Macro F1': probability_f1,
            '多數類基準正確率': majority_accuracy,
            '相對基準差': ((probability_accuracy if np.isfinite(probability_accuracy)
                           else point_accuracy) - majority_accuracy),
            '平衡正確率': balanced,
            'Macro F1': float(f1_score(y, pred, labels=[0, 1, 2], average='macro', zero_division=0)),
            '下行先觸召回率': float(recalls[0]), '盤整召回率': float(recalls[1]),
            '上行先觸召回率': float(recalls[2]), '10日預測報酬MAE': mae,
            '10日預測報酬RMSE': rmse, '10日預測報酬偏誤': bias,
            '10日方向命中率': direction, '預測與實際報酬Spearman': rho,
            'Log Loss': logloss, '基準Log Loss': base_logloss,
            'Brier': brier, '基準Brier': base_brier, 'Brier Skill': skill,
            '邊際校準誤差': marginal_error,
            '判定': verdict,
        })
    return pd.DataFrame(rows, columns=columns)


def _model_comparison_table(ai, timesfm):
    """Compare models against the same realised path without hiding probability origin."""
    columns = ['市場', '模型', '類型', '成熟樣本', '正確率', '平衡正確率', 'Macro F1',
               'Log Loss', '基準Log Loss', 'Brier', '基準Brier', 'Brier Skill',
               '邊際校準誤差', '判定']
    rows = []
    for _, row in ai.iterrows():
        rows.append({name: row.get(name) for name in columns} | {'類型': '三分類AI原生機率'})
    for _, row in timesfm.iterrows():
        rows.append({
            '市場': row.get('市場'), '模型': row.get('模型'),
            '類型': 'TimesFM分位情境估計機率', '成熟樣本': row.get('機率成熟樣本'),
            '正確率': row.get('機率最高類正確率'),
            '平衡正確率': row.get('機率平衡正確率'), 'Macro F1': row.get('機率Macro F1'),
            'Log Loss': row.get('Log Loss'), '基準Log Loss': row.get('基準Log Loss'),
            'Brier': row.get('Brier'), '基準Brier': row.get('基準Brier'),
            'Brier Skill': row.get('Brier Skill'), '邊際校準誤差': row.get('邊際校準誤差'),
            '判定': row.get('判定'),
        })
    return pd.DataFrame(rows, columns=columns)


POSITIVE = ('買進', '強買', '偏多', '多頭', '強勢', '偏強', '正向', '安全', '優良', '增加', '成長', '低')
NEGATIVE = ('賣出', '減碼', '偏空', '空頭', '轉弱', '弱勢', '偏弱', '危險', '過熱', '衰退', '高')
CATEGORY_HINTS = ('短期趨勢', '中期趨勢', '長期趨勢', '籌碼', '量價參考', '市場環境', '市場情緒',
                  '個股情緒', '新聞情緒', '融資壓力', '風險',
                  '未持有建議', '基本面投資建議', '盈餘投資建議',
                  '護城河', '財務安全性', '獲利品質', '相對大盤動能')
NUMBER_HINTS = ('歷史PE百分位_1Y', '歷史PE百分位_3Y', '歷史PE百分位_5Y', '成長分數', '品質分數', '巴菲特檢核分數', '基本面分數',
                'RSI', 'ATR比例', 'CMF', '籌碼K線分數', '投入資本報酬率', '股東權益報酬率',
                '自由現金流殖利率', '市場情緒分數', '個股情緒分數', '新聞情緒分數',
                '融資限額使用率', '融資餘額近N日變化率', '大盤融資近5日變化率',
                '大盤融資近20日變化率', 'VIX', 'SPY近20日報酬率', '相對SPY20日報酬差',
                '相對大盤20日報酬差', '相對大盤60日報酬差', '相對大盤120日報酬差',
                '60日風險調整相對動能', 'TimesFM估計上行先觸機率',
                'TimesFM估計下行先觸機率', 'TimesFM估計盤整機率')


def _expected_sign(value):
    text = str(value)
    if any(token in text for token in POSITIVE):
        return 1
    if any(token in text for token in NEGATIVE):
        return -1
    return 0


def _evaluation_target(column):
    name = str(column)
    if name.startswith('WHID_'):
        return 'return_126', '126交易日'
    if '市場情緒' in name:
        return 'benchmark_return_10', '市場10交易日'
    if '相對大盤' in name or '風險調整相對動能' in name:
        return 'abnormal_return_10', '相對市場10交易日'
    if any(token in name for token in ['個股情緒','新聞情緒','融資壓力','融資限額使用率','融資餘額近N日變化率']):
        return 'abnormal_return_10', '相對市場10交易日'
    if '短期趨勢' in name:
        return 'return_5', '5交易日'
    if '長期趨勢' in name:
        return 'return_60', '60交易日'
    return 'return_10', '10交易日'


def _rolling_slope(values):
    values = np.asarray(values, dtype=float)
    valid = np.isfinite(values)
    if valid.sum() < 3:
        return np.nan
    x = np.arange(len(values), dtype=float)[valid]
    return float(np.polyfit(x, values[valid], 1)[0])


def _last_percentile(values):
    values = np.asarray(values, dtype=float)
    if not len(values) or not np.isfinite(values[-1]):
        return np.nan
    valid = values[np.isfinite(values)]
    return float((valid <= values[-1]).mean()) if len(valid) >= 3 else np.nan


def _training_dataset(frame):
    """Build an as-of feature table and past-only trend/relative features."""
    if frame.empty:
        return pd.DataFrame(), [], [], []
    snapshots = pd.DataFrame([json.loads(value) for value in frame.snapshot_json]).reset_index(drop=True)
    outcomes = frame[['market', 'ticker', 'signal_date', 'actual_class',
                      'return_5', 'return_10', 'return_20', 'return_60', 'return_126', 'return_252',
                      'benchmark_return_5', 'benchmark_return_10', 'benchmark_return_20',
                      'abnormal_return_5', 'abnormal_return_10', 'abnormal_return_20',
                      'max_up_10', 'max_down_10']].reset_index(drop=True)
    category_columns = [column for column in snapshots
                        if any(hint in str(column) for hint in CATEGORY_HINTS)
                        and not pd.api.types.is_numeric_dtype(snapshots[column])]
    number_columns = [column for column in snapshots
                      if any(hint in str(column) for hint in NUMBER_HINTS)]
    metadata_columns = [column for column in ['WHID_產業', '產業'] if column in snapshots]
    selected = list(dict.fromkeys(metadata_columns + category_columns + number_columns))
    data = pd.concat([outcomes, snapshots.reindex(columns=selected)], axis=1)
    data['signal_date'] = pd.to_datetime(data.signal_date, errors='coerce')
    data = data.sort_values(['market', 'ticker', 'signal_date']).reset_index(drop=True)
    derived = []
    for column in number_columns:
        data[column] = pd.to_numeric(data[column], errors='coerce')
        change = f'變化1次｜{column}'
        slope = f'斜率5次｜{column}'
        acceleration = f'加速度5次｜{column}'
        own_rank = f'自身百分位20次｜{column}'
        market_rank = f'市場相對百分位｜{column}'
        grouped = data.groupby(['market', 'ticker'], sort=False)[column]
        data[change] = grouped.diff()
        data[slope] = grouped.transform(
            lambda series: series.rolling(5, min_periods=3).apply(_rolling_slope, raw=True))
        data[acceleration] = data.groupby(['market', 'ticker'], sort=False)[slope].diff()
        data[own_rank] = grouped.transform(
            lambda series: series.rolling(20, min_periods=3).apply(_last_percentile, raw=True))
        data[market_rank] = data.groupby(['market', 'signal_date'])[column].rank(pct=True)
        derived.extend([change, slope, acceleration, own_rank, market_rank])
        industry_column = 'WHID_產業' if 'WHID_產業' in data else ('產業' if '產業' in data else None)
        if industry_column:
            industry_rank = f'產業相對百分位｜{column}'
            industry_group = data.groupby(['market', 'signal_date', industry_column])[column]
            data[industry_rank] = industry_group.rank(pct=True)
            data.loc[industry_group.transform('count') < 3, industry_rank] = np.nan
            derived.append(industry_rank)
    data['signal_date'] = data.signal_date.dt.strftime('%Y-%m-%d')
    return data, number_columns, derived, category_columns


def _feature_target(feature):
    original = str(feature).split('｜', 1)[-1]
    return _evaluation_target(original)


def _derived_numeric_table(training, features):
    if training.empty or not features:
        return pd.DataFrame()
    from scipy.stats import spearmanr
    rows = []
    for market, market_frame in training.groupby('market'):
        for feature in features:
            target, target_label = _feature_target(feature)
            selected = pd.DataFrame({
                'value': pd.to_numeric(market_frame[feature], errors='coerce'),
                'future_return': pd.to_numeric(market_frame[target], errors='coerce'),
            }).dropna()
            kind = str(feature).split('｜', 1)[0]
            if len(selected) < 10 or selected.value.nunique() < 3:
                rows.append({'市場': market, '衍生特徵': feature, '特徵類型': kind,
                             '評估期間': target_label, '成熟樣本': len(selected),
                             '判定': '資料不足'})
                continue
            rho, pvalue = spearmanr(selected.value, selected.future_return)
            low = selected.value <= selected.value.quantile(.25)
            high = selected.value >= selected.value.quantile(.75)
            spread = selected.loc[high, 'future_return'].mean() - selected.loc[low, 'future_return'].mean()
            verdict = '資料不足' if len(selected) < 30 else (
                '初步有關聯' if abs(rho) >= .15 and pvalue <= .10 else
                '關聯偏弱／需修正' if abs(rho) < .08 else '尚不明確')
            rows.append({
                '市場': market, '衍生特徵': feature, '特徵類型': kind,
                '評估期間': target_label, '成熟樣本': len(selected),
                'Spearman相關': rho, 'p值': pvalue,
                '高低四分位報酬差': spread,
                '低值組上漲比例': (selected.loc[low, 'future_return'] > 0).mean(),
                '高值組上漲比例': (selected.loc[high, 'future_return'] > 0).mean(),
                '判定': verdict,
            })
    return pd.DataFrame(rows)


def _combination_series(training):
    combinations = []

    def signs(column):
        return training[column].map(_expected_sign) if column in training else None

    trend_columns = ['短期趨勢', '中期趨勢', '長期趨勢']
    if all(column in training for column in trend_columns):
        trend_signs = pd.concat([signs(column) for column in trend_columns], axis=1)
        value = pd.Series('周期分歧', index=training.index)
        value.loc[(trend_signs > 0).all(axis=1)] = '偏多共振'
        value.loc[(trend_signs < 0).all(axis=1)] = '偏空共振'
        combinations.append(('短中長趨勢組合', value, 'return_10', '10交易日'))

    if '中期趨勢' in training and ('籌碼' in training or '量價參考' in training):
        flow = pd.Series(0, index=training.index)
        tw = training.market == 'TW'
        us = training.market == 'US'
        if '籌碼' in training:
            flow.loc[tw] = signs('籌碼').loc[tw]
        if '量價參考' in training:
            flow.loc[us] = signs('量價參考').loc[us]
        trend = signs('中期趨勢')
        value = pd.Series('方向不一致', index=training.index)
        value.loc[(trend > 0) & (flow > 0)] = '偏多共振'
        value.loc[(trend < 0) & (flow < 0)] = '偏空共振'
        combinations.append(('中期趨勢＋籌碼量價', value, 'return_10', '10交易日'))

    pairs = [
        ('基本面＋盈餘品質', 'WHID_基本面投資建議', 'WHID_盈餘投資建議'),
        ('現金品質＋財務安全', 'WHID_現金獲利品質評價', 'WHID_財務安全性評價'),
    ]
    for name, left, right in pairs:
        if left not in training or right not in training:
            continue
        left_sign, right_sign = signs(left), signs(right)
        value = pd.Series('方向不一致', index=training.index)
        value.loc[(left_sign > 0) & (right_sign > 0)] = '偏多共振'
        value.loc[(left_sign < 0) & (right_sign < 0)] = '偏空共振'
        combinations.append((name, value, 'return_126', '126交易日'))
    return combinations


def _combination_table(training):
    rows = []
    for name, values, target, target_label in _combination_series(training):
        for (market, value), indexes in pd.DataFrame({
                'market': training.market, 'value': values}).groupby(['market', 'value']).groups.items():
            selected = training.loc[list(indexes)].dropna(subset=[target])
            sign = 1 if '偏多' in value else -1 if '偏空' in value else 0
            returns = pd.to_numeric(selected[target], errors='coerce').dropna()
            directional = returns * sign if sign else pd.Series(dtype=float)
            rows.append({
                '市場': market, '組合': name, '組合狀態': value, '評估期間': target_label,
                '成熟樣本': len(returns), '平均報酬': returns.mean(),
                '上漲比例': (returns > 0).mean() if len(returns) else np.nan,
                '方向命中率': (directional > 0).mean() if len(directional) else np.nan,
                '判定': ('資料不足' if len(returns) < 30 else
                         '初步有決策價值' if len(directional) and (directional > 0).mean() >= .55 else
                         '方向需修正' if len(directional) and (directional > 0).mean() <= .45 else
                         '尚不明確'),
            })
    return pd.DataFrame(rows)


def _decision_value_table(training, category_columns):
    rows = []
    if training.empty:
        return pd.DataFrame()
    for market, market_frame in training.groupby('market'):
        for column in category_columns:
            target, target_label = _evaluation_target(column)
            values = market_frame[column]
            signs = values.map(_expected_sign)
            future = pd.to_numeric(market_frame[target], errors='coerce')
            valid = values.notna() & future.notna()
            actionable = valid & (signs != 0)
            if not actionable.any():
                continue
            directional = future.loc[actionable] * signs.loc[actionable]
            downside = np.nan
            if target == 'return_10':
                adverse = np.where(signs.loc[actionable] > 0,
                                   market_frame.loc[actionable, 'max_down_10'],
                                   -market_frame.loc[actionable, 'max_up_10'])
                downside = pd.to_numeric(pd.Series(adverse), errors='coerce').mean()
            hit_rate = (directional > 0).mean()
            rows.append({
                '市場': market, '指標': column, '評估期間': target_label,
                '有效樣本': int(valid.sum()), '可採取方向樣本': int(actionable.sum()),
                '訊號覆蓋率': actionable.sum() / valid.sum(),
                '順訊號平均報酬': directional.mean(), '順訊號中位報酬': directional.median(),
                '方向命中率': hit_rate, '順訊號報酬10分位': directional.quantile(.10),
                '10日平均不利波動': downside,
                '判定': ('資料不足' if actionable.sum() < 30 else
                         '初步可作決策參考' if hit_rate >= .55 and directional.mean() > 0 else
                         '方向需修正' if hit_rate <= .45 else '尚不明確'),
            })
    return pd.DataFrame(rows)


def _stability_table(training, features):
    if training.empty:
        return pd.DataFrame()
    from scipy.stats import spearmanr
    work = training.copy()
    work['month'] = pd.to_datetime(work.signal_date).dt.strftime('%Y-%m')
    rows = []
    for market, market_frame in work.groupby('market'):
        for feature in features:
            target, target_label = _feature_target(feature)
            overall = pd.DataFrame({
                'value': pd.to_numeric(market_frame[feature], errors='coerce'),
                'future': pd.to_numeric(market_frame[target], errors='coerce'),
            }).dropna()
            if len(overall) < 10 or overall.value.nunique() < 3:
                continue
            overall_rho = float(spearmanr(overall.value, overall.future).statistic)
            reference_direction = np.sign(overall_rho)

            def grouped_correlations(column):
                values = []
                if column is None or column not in market_frame:
                    return values
                for label, group in market_frame.groupby(column, dropna=True):
                    selected = pd.DataFrame({
                        'value': pd.to_numeric(group[feature], errors='coerce'),
                        'future': pd.to_numeric(group[target], errors='coerce'),
                    }).dropna()
                    if len(selected) < 10 or selected.value.nunique() < 3:
                        continue
                    values.append((str(label), len(selected),
                                   float(spearmanr(selected.value, selected.future).statistic)))
                return values

            monthly = grouped_correlations('month')
            regime_column = next((column for column in ['市場情緒', '市場環境']
                                  if column in market_frame and market_frame[column].notna().any()), None)
            regimes = grouped_correlations(regime_column)
            industry_column = next((column for column in ['WHID_產業', '產業']
                                    if column in market_frame and market_frame[column].notna().any()), None)
            industries = grouped_correlations(industry_column)

            def summary(values):
                correlations = pd.Series([item[2] for item in values], dtype=float)
                if correlations.empty:
                    return np.nan, np.nan, np.nan, np.nan
                consistency = ((np.sign(correlations) == reference_direction).mean()
                               if reference_direction else np.nan)
                return correlations.median(), correlations.min(), correlations.max(), consistency

            month_median, month_min, month_max, month_consistency = summary(monthly)
            regime_median, _, _, regime_consistency = summary(regimes)
            industry_median, _, _, industry_consistency = summary(industries)
            verdict = '需至少2個可評估月份'
            if len(monthly) >= 2:
                if month_consistency < .70 or abs(overall_rho) < .10:
                    verdict = '跨月不穩定／需修正'
                elif len(regimes) >= 2 and regime_consistency < .60:
                    verdict = '市場狀態下不穩定／需修正'
                elif len(industries) >= 2 and industry_consistency < .60:
                    verdict = '產業間不穩定／需修正'
                else:
                    verdict = '初步穩定'
            rows.append({
                '市場': market, '特徵': feature, '評估期間': target_label,
                '整體Spearman': overall_rho,
                '可評估月份': len(monthly),
                '月份範圍': f'{monthly[0][0]}～{monthly[-1][0]}' if monthly else None,
                '月相關中位數': month_median, '月相關最小值': month_min,
                '月相關最大值': month_max, '方向一致月份比例': month_consistency,
                '市場狀態欄位': regime_column, '可評估市場狀態數': len(regimes),
                '市場狀態相關中位數': regime_median, '市場狀態方向一致比例': regime_consistency,
                '產業欄位': industry_column, '可評估產業數': len(industries),
                '產業相關中位數': industry_median, '產業方向一致比例': industry_consistency,
                '判定': verdict,
            })
    return pd.DataFrame(rows)


def _walk_forward_table(training, features, embargo_dates=10):
    """Experimental multivariate screen using an embargoed chronological holdout."""
    columns = ['市場', '狀態', '訓練樣本', '測試樣本', '訓練截止日', '測試起始日',
               '隔離日期數', '特徵數', '使用特徵', '正確率', '多數類基準正確率',
               '平衡正確率', 'Macro F1', 'Log Loss', 'Brier', '判定']
    if training.empty:
        return pd.DataFrame(columns=columns)
    from sklearn.impute import SimpleImputer
    from sklearn.linear_model import LogisticRegression
    from sklearn.metrics import accuracy_score, balanced_accuracy_score, f1_score, log_loss
    from sklearn.pipeline import make_pipeline
    from sklearn.preprocessing import StandardScaler

    label_map = {name: index for index, name in enumerate(CLASS_ORDER)}
    rows = []
    for market, group in training.groupby('market'):
        group = group.loc[group.actual_class.isin(CLASS_ORDER)].copy()
        dates = sorted(pd.to_datetime(group.signal_date).dropna().unique())
        if len(group) < 120 or len(dates) < 30:
            rows.append({'市場': market, '狀態': '資料不足', '訓練樣本': 0,
                         '測試樣本': 0, '隔離日期數': embargo_dates,
                         '判定': '至少需120筆成熟樣本及30個不同訊號日'})
            continue
        split = max(int(len(dates) * .80), embargo_dates + 10)
        test_dates = dates[split:]
        train_dates = dates[:max(0, split - embargo_dates)]
        train = group.loc[pd.to_datetime(group.signal_date).isin(train_dates)].copy()
        test = group.loc[pd.to_datetime(group.signal_date).isin(test_dates)].copy()
        candidates = []
        y_train = train.actual_class.map(label_map)
        for feature in features:
            values = pd.to_numeric(train[feature], errors='coerce')
            selected = pd.DataFrame({'value': values, 'target': y_train}).dropna()
            if len(selected) >= 30 and selected.value.nunique() >= 3:
                correlation = selected.value.corr(selected.target, method='spearman')
                candidates.append((abs(correlation) if np.isfinite(correlation) else 0, feature))
        chosen = [feature for _, feature in sorted(candidates, reverse=True)[:20]]
        if len(train) < 60 or len(test) < 20 or y_train.nunique() < 2 or not chosen:
            rows.append({'市場': market, '狀態': '資料不足', '訓練樣本': len(train),
                         '測試樣本': len(test), '隔離日期數': embargo_dates,
                         '判定': '時間切割後的訓練、測試或類別不足'})
            continue
        model = make_pipeline(SimpleImputer(strategy='median'), StandardScaler(),
                              LogisticRegression(max_iter=1000, class_weight='balanced'))
        model.fit(train[chosen], y_train)
        raw_probability = model.predict_proba(test[chosen])
        probability = np.full((len(test), 3), 1e-9)
        probability[:, model[-1].classes_.astype(int)] = raw_probability
        probability /= probability.sum(axis=1, keepdims=True)
        y_test = test.actual_class.map(label_map).to_numpy(dtype=int)
        pred = probability.argmax(axis=1)
        train_distribution = np.bincount(y_train, minlength=3) / len(y_train)
        train_majority_class = int(train_distribution.argmax())
        baseline = float((y_test == train_majority_class).mean())
        one_hot = np.eye(3)[y_test]
        brier = float(np.mean(np.sum((probability - one_hot) ** 2, axis=1)))
        accuracy = float(accuracy_score(y_test, pred))
        rows.append({
            '市場': market, '狀態': '完成時間序列樣本外驗證',
            '訓練樣本': len(train), '測試樣本': len(test),
            '訓練截止日': str(pd.Timestamp(train_dates[-1]).date()),
            '測試起始日': str(pd.Timestamp(test_dates[0]).date()),
            '隔離日期數': embargo_dates, '特徵數': len(chosen), '使用特徵': '；'.join(chosen),
            '正確率': accuracy, '多數類基準正確率': baseline,
            '平衡正確率': float(balanced_accuracy_score(y_test, pred)),
            'Macro F1': float(f1_score(y_test, pred, labels=[0, 1, 2], average='macro', zero_division=0)),
            'Log Loss': float(log_loss(y_test, probability, labels=[0, 1, 2])),
            'Brier': brier,
            '判定': '初步有增量價值' if accuracy > baseline else '未優於多數類基準',
        })
    return pd.DataFrame(rows, columns=columns)


def _indicator_tables(frame, training=None):
    if frame.empty:
        return pd.DataFrame(), pd.DataFrame()
    from scipy.stats import spearmanr
    if training is None:
        training, _, _, _ = _training_dataset(frame)
    categorical_rows, numeric_rows = [], []
    category_columns = [column for column in training
                        if any(hint in str(column) for hint in CATEGORY_HINTS)
                        and not pd.api.types.is_numeric_dtype(training[column])]
    number_columns = [column for column in training
                      if any(hint in str(column) for hint in NUMBER_HINTS)
                      and '｜' not in str(column)]
    for market, market_frame in training.groupby('market'):
        for column in category_columns:
            values = market_frame[column]
            target, target_label = _evaluation_target(column)
            for value, indexes in values.groupby(values, dropna=True).groups.items():
                selected = market_frame.loc[list(indexes)].dropna(subset=[target])
                if selected.empty:
                    categorical_rows.append({
                        '市場': market, '指標': column, '分類值': value, '評估期間': target_label,
                        '預期方向': '尚未判定', '成熟樣本': 0, '判定': f'需累積{target_label}結果',
                    })
                    continue
                sign = _expected_sign(value)
                returns = selected[target]
                hit = ((returns > 0) if sign > 0 else
                       (returns < 0) if sign < 0 else pd.Series(dtype=bool))
                hit_rate = float(hit.mean()) if sign else np.nan
                verdict = '僅描述／待定義方向'
                if sign and len(selected) < 30:
                    verdict = '資料不足'
                elif sign and hit_rate >= .55:
                    verdict = '初步可信'
                elif sign and hit_rate <= .45:
                    verdict = '方向相反；需修正'
                elif sign:
                    verdict = '尚不明確'
                categorical_rows.append({
                    '市場': market, '指標': column, '分類值': value, '評估期間': target_label,
                    '預期方向': '正向' if sign > 0 else '負向' if sign < 0 else '未定義',
                    '成熟樣本': len(selected), '期間平均報酬': returns.mean(),
                    '期間上漲比例': (returns > 0).mean(),
                    '方向命中率': hit_rate, '上行先觸率': (selected.actual_class == 'up_first').mean(),
                    '下行先觸率': (selected.actual_class == 'down_first').mean(), '判定': verdict,
                })
        for column in number_columns:
            values = pd.to_numeric(market_frame[column], errors='coerce')
            target, target_label = _evaluation_target(column)
            selected = pd.DataFrame({'value': values, 'future_return': market_frame[target]}).dropna()
            if len(selected) < 10 or selected.value.nunique() < 3:
                numeric_rows.append({'市場': market, '指標': column, '評估期間': target_label,
                                     '成熟樣本': len(selected), '判定': f'需累積{target_label}結果'})
                continue
            rho, pvalue = spearmanr(selected.value, selected.future_return)
            verdict = '資料不足' if len(selected) < 30 else (
                '初步有關聯' if abs(rho) >= .15 and pvalue <= .10 else
                '關聯偏弱／需修正' if abs(rho) < .08 else '尚不明確')
            numeric_rows.append({
                '市場': market, '指標': column, '評估期間': target_label,
                '成熟樣本': len(selected), 'Spearman相關': rho,
                'p值': pvalue,
                '低值組期間平均報酬': selected.loc[selected.value <= selected.value.quantile(.25), 'future_return'].mean(),
                '高值組期間平均報酬': selected.loc[selected.value >= selected.value.quantile(.75), 'future_return'].mean(),
                '判定': verdict,
            })
    return pd.DataFrame(categorical_rows), pd.DataFrame(numeric_rows)


def review_output_path(now=None, root=REVIEW_ROOT):
    current = now or datetime.now()
    folder = Path(root) / current.strftime('%Y') / current.strftime('%m')
    return folder / f"WHID_30日驗證_{current.strftime('%Y%m%d_%H%M%S')}.xlsx"


def build_review(db_path=DB_PATH, minimum_calendar_days=30, output_root=REVIEW_ROOT, now=None):
    connection = connect(db_path)
    forecasts = pd.read_sql_query(
        "SELECT * FROM forecasts WHERE settlement_status='settled'", connection)
    indicators = pd.read_sql_query(
        "SELECT * FROM indicators WHERE settlement_status='settled'", connection)
    coverage = pd.read_sql_query('''SELECT market, MIN(signal_date) AS first_signal_date,
        MAX(signal_date) AS last_signal_date, COUNT(*) AS snapshots,
        SUM(CASE WHEN settlement_status='settled' THEN 1 ELSE 0 END) AS settled
        FROM indicators GROUP BY market''', connection)
    connection.close()
    ai, calibration = _ai_tables(forecasts)
    timesfm = _timesfm_table(indicators)
    model_comparison = _model_comparison_table(ai, timesfm)
    training, raw_features, derived_features, category_features = _training_dataset(indicators)
    categorical, numeric = _indicator_tables(indicators, training)
    derived = _derived_numeric_table(training, derived_features)
    combinations = _combination_table(training)
    decision_value = _decision_value_table(training, category_features)
    stability = _stability_table(training, raw_features + derived_features)
    walk_forward = _walk_forward_table(training, raw_features + derived_features)
    span_days = 0
    if not coverage.empty:
        span_days = max((pd.to_datetime(coverage.last_signal_date) - pd.to_datetime(coverage.first_signal_date)).dt.days.max(), 0)
    readiness = '已達30日，可進行第一輪判讀' if span_days >= minimum_calendar_days else f'累積{span_days}日，尚未達{minimum_calendar_days}日'
    overview = pd.DataFrame([
        {'項目': '資料狀態', '結果': readiness},
        {'項目': 'AI成熟預測數', '結果': len(forecasts)},
        {'項目': 'TimesFM成熟預測數', '結果': int(timesfm['成熟樣本'].sum()) if not timesfm.empty else 0},
        {'項目': '指標成熟快照數', '結果': len(indicators)},
        {'項目': '判讀原則', '結果': '至少30筆成熟樣本；AI須優於類別基準，方向型指標命中率以55%為初步門檻'},
        {'項目': '期限分工', '結果': '短期趨勢5日；中期／籌碼／量價／風險10日；長期趨勢60日；WHID基本面與估值126日'},
        {'項目': '進階分析', '結果': '包含指標變化、5次斜率與加速度、自身20次百分位、市場及產業相對百分位、組合效果、跨月穩定性、決策價值與隔離式時間序列樣本外驗證'},
        {'項目': '限制', '結果': '30日只足以初查第二階段；第一階段須等待約6個月，不可用短期漲跌判定'},
    ])
    method = pd.DataFrame([
        {'主題': '封存', '說明': '同一市場、股票、訊號日與模型只保留首次預測，後續重跑不覆寫。'},
        {'主題': '成熟', '說明': '訊號日後第10個交易日完成才核對；同日上下雙觸採下行先觸。'},
        {'主題': 'AI', '說明': '比較正確率、平衡正確率、Macro F1、Log Loss、Brier Skill與校準誤差。'},
        {'主題': 'TimesFM', '說明': '以0.1至0.9共9條分位路徑的等權情境占比估算三分類機率，使用同一真實10日ATR路徑計算Log Loss、Brier與校準；此為近似值，不是原生分類機率，也不併入AI集成。'},
        {'主題': '分類指標', '說明': '短期趨勢看5日，中期／籌碼／風險看10日，長期趨勢看60日，WHID基本面與估值看126日。'},
        {'主題': '數值指標', '說明': '依指標期限使用Spearman相關及高低四分位報酬差；多重比較結果只作篩選。'},
        {'主題': '趨勢與相對值', '說明': '所有衍生值只使用當日及過去已封存快照：前次變化、最近5次斜率與加速度、自身最近20次百分位、同市場及同產業同日橫斷面百分位；產業少於3檔不計排名。'},
        {'主題': '組合與決策', '說明': '分別檢查趨勢共振、趨勢加籌碼量價、基本面加盈餘品質，以及訊號覆蓋率、順訊號報酬、命中率和不利波動。'},
        {'主題': '樣本外驗證', '說明': '依日期切割訓練與測試，中間隔離10個訊號日；至少120筆成熟樣本及30個不同訊號日才執行，特徵篩選僅使用訓練區。'},
        {'主題': '訓練資料', '說明': '訓練資料工作表保存當時可見的原始與衍生特徵及事後標籤；不得使用尚未成熟結果，也不得隨機切割時間序列。'},
    ])
    output = review_output_path(now=now, root=output_root)
    output.parent.mkdir(parents=True, exist_ok=True)
    sheets = {
        '總覽': overview, '資料涵蓋': coverage, '模型比較': model_comparison,
        'AI模型': ai, 'AI校準': calibration, 'TimesFM驗證': timesfm,
        '分類指標': categorical, '數值指標': numeric, '衍生指標': derived,
        '組合效果': combinations, '穩定性': stability, '決策價值': decision_value,
        '樣本外驗證': walk_forward, '訓練資料': training, '方法': method,
    }
    from openpyxl.styles import Font, PatternFill
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        for name, frame in sheets.items():
            (frame if not frame.empty else pd.DataFrame({'狀態': ['尚無成熟資料']})).to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(color='FFFFFF', bold=True)
                cell.fill = PatternFill('solid', fgColor='17365D')
            for column in ws.columns:
                width = min(55, max(12, max(len(str(cell.value or '')) for cell in column) + 2))
                ws.column_dimensions[column[0].column_letter].width = width
    return {'path': str(output), 'readiness': readiness, 'ai_rows': len(ai),
            'timesfm_rows': len(timesfm), 'model_comparison_rows': len(model_comparison),
            'categorical_rows': len(categorical), 'numeric_rows': len(numeric),
            'derived_rows': len(derived), 'combination_rows': len(combinations),
            'stability_rows': len(stability), 'decision_rows': len(decision_value),
            'training_rows': len(training)}


def status(db_path=DB_PATH):
    connection = connect(db_path)
    tables = {}
    for table in ['forecasts', 'indicators']:
        tables[table] = pd.read_sql_query(
            f'''SELECT market, COUNT(*) AS total,
            SUM(CASE WHEN settlement_status='settled' THEN 1 ELSE 0 END) AS settled,
            MIN(signal_date) AS first_date, MAX(signal_date) AS last_date
            FROM {table} GROUP BY market''', connection).to_dict('records')
    connection.close()
    return tables


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('action', choices=['archive', 'settle', 'review', 'status'])
    parser.add_argument('--market', choices=['TW', 'US'])
    parser.add_argument('--report')
    parser.add_argument('--db', default=str(DB_PATH))
    args = parser.parse_args()
    if args.action == 'archive':
        if not args.market:
            parser.error('archive需要--market')
        result = archive_report(args.market, args.report, args.db) if args.report else archive_latest(args.market, db_path=args.db)
    elif args.action == 'settle':
        result = settle_pending(args.db)
    elif args.action == 'review':
        result = build_review(args.db)
    else:
        result = status(args.db)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
